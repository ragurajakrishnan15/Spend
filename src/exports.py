"""Stakeholder exports: an Excel workbook (live formulas, lookups, chart) and a PowerPoint deck.

Both are built from the same computed context, so every number is derived from the data - none are
hard-coded. They are plain openpyxl / python-pptx files; no proprietary add-ins are involved.
"""
from __future__ import annotations

import io
from datetime import datetime

import pandas as pd

from src.analytics import (category_spend, customer_rfm, forecast_merchant_spend, forecast_monthly_spend,
                           monthly_spend, segment_summary)

DISCLAIMER = ("Synthetic data only. Independent educational project - not affiliated with any bank; "
              "no claim of regulatory compliance. Not for credit, eligibility or other decisions about people.")


def build_context(df: pd.DataFrame, label: str = "All transactions", horizon: int = 6) -> dict:
    """Compute everything the exports need from a (possibly filtered) transaction frame."""
    if df.empty:
        raise ValueError("No transactions match the selected filters.")
    monthly = monthly_spend(df)
    fc = forecast_monthly_spend(monthly, horizon=horizon)
    seg = customer_rfm(df)
    merchants = forecast_merchant_spend(df) if "merchant_clean" in df else {"error": "No merchant data."}
    return {
        "label": label,
        "generated": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "period": f"{df['timestamp'].min():%Y-%m-%d} to {df['timestamp'].max():%Y-%m-%d}",
        "kpis": {
            "total_spend": float(df["amount"].sum()), "transactions": int(len(df)),
            "customers": int(df["customer_id"].nunique()), "avg_transaction": float(df["amount"].mean()),
            "response_rate": float(df["campaign_response"].mean()),
        },
        "monthly": monthly,
        "categories": category_spend(df),
        "forecast": fc,
        "segments": segment_summary(seg),
        "merchants": merchants,
    }


def _takeaways(ctx: dict) -> list[str]:
    m = ctx["monthly"].set_index("month")["spend"]
    out = []
    if len(m):
        out.append(f"Peak month: {m.idxmax():%b %Y} at ${m.max():,.0f}; lowest: {m.idxmin():%b %Y} at ${m.min():,.0f}.")
    if len(m) >= 24:
        y1, y0 = m.iloc[-12:].sum(), m.iloc[-24:-12].sum()
        out.append(f"Trailing 12 months spend ${y1:,.0f}, {((y1 / y0) - 1) * 100:+.1f}% vs the prior 12 months.")
    cat = ctx["categories"]
    if len(cat):
        top = cat.iloc[0]
        out.append(f"{top['category']} is the largest category at {top['spend'] / cat['spend'].sum() * 100:.0f}% of spend.")
    fc = ctx["forecast"]
    if fc["metrics"]:
        out.append(f"Best backtest model: {fc['best_model']} (lowest MAE over {fc['holdout_months']} held-out months).")
    return out


# ------------------------------------------------------------------------------------------ Excel
def build_excel(ctx: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    navy = PatternFill("solid", fgColor="0B1F3A")
    head_font = Font(bold=True, color="FFFFFF")
    wb = Workbook()

    def header(ws, row, labels):
        for c, t in enumerate(labels, 1):
            cell = ws.cell(row=row, column=c, value=t)
            cell.fill, cell.font = navy, head_font
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

    # Monthly (data + formulas)
    wm = wb.active
    wm.title = "Monthly"
    header(wm, 1, ["Month", "Spend", "MoM %", "YoY %", "3-month avg"])
    monthly = ctx["monthly"]
    for r, (month, spend) in enumerate(zip(monthly["month"], monthly["spend"]), 2):
        wm.cell(r, 1, month.to_pydatetime()).number_format = "mmm yyyy"
        wm.cell(r, 2, round(float(spend), 2)).number_format = '"$"#,##0'
        if r > 2:
            wm.cell(r, 3, f"=B{r}/B{r-1}-1").number_format = "0.0%"
        if r > 13:
            wm.cell(r, 4, f"=B{r}/B{r-12}-1").number_format = "0.0%"
        if r > 3:
            wm.cell(r, 5, f"=AVERAGE(B{r-2}:B{r})").number_format = '"$"#,##0'
    last_m = len(monthly) + 1
    wm.column_dimensions["A"].width = 14
    for col in "BCDE":
        wm.column_dimensions[col].width = 15
    chart = LineChart()
    chart.title, chart.height, chart.width = "Monthly spend", 8, 18
    chart.add_data(Reference(wm, min_col=2, min_row=1, max_row=last_m), titles_from_data=True)
    chart.set_categories(Reference(wm, min_col=1, min_row=2, max_row=last_m))
    wm.add_chart(chart, "G2")

    # Categories
    wc = wb.create_sheet("Categories")
    header(wc, 1, ["Category", "Spend", "Share of spend"])
    cats = ctx["categories"].reset_index(drop=True)
    for r, (c, s) in enumerate(zip(cats["category"], cats["spend"]), 2):
        wc.cell(r, 1, c)
        wc.cell(r, 2, round(float(s), 2)).number_format = '"$"#,##0'
        wc.cell(r, 3, f"=B{r}/SUM($B$2:$B${len(cats)+1})").number_format = "0.0%"
    wc.column_dimensions["A"].width = 18
    wc.column_dimensions["B"].width = 16
    wc.column_dimensions["C"].width = 16
    cat_last = len(cats) + 1

    # Summary with INDEX/MATCH and VLOOKUP tools
    ws = wb.create_sheet("Summary", 0)
    ws["A1"], ws["A1"].font = "SpendIQ - Spend summary", Font(bold=True, size=16, color="0B1F3A")
    ws["A2"] = f"Scope: {ctx['label']} | Period: {ctx['period']} | Generated {ctx['generated']}"
    ws["A3"], ws["A3"].font = DISCLAIMER, Font(italic=True, color="7A1F1F")
    k = ctx["kpis"]
    ws["A5"], ws["B5"] = "Total spend (sum of Monthly)", f"=SUM(Monthly!B2:B{last_m})"
    ws["A6"], ws["B6"] = "Average monthly spend", f"=AVERAGE(Monthly!B2:B{last_m})"
    ws["A7"], ws["B7"] = "Latest month spend", f"=INDEX(Monthly!B2:B{last_m},MATCH(MAX(Monthly!A2:A{last_m}),Monthly!A2:A{last_m},0))"
    ws["A8"], ws["B8"] = "Latest month MoM %", f"=INDEX(Monthly!C2:C{last_m},MATCH(MAX(Monthly!A2:A{last_m}),Monthly!A2:A{last_m},0))"
    ws["A9"], ws["B9"] = "Transactions", k["transactions"]
    ws["A10"], ws["B10"] = "Customers", k["customers"]
    ws["A11"], ws["B11"] = "Average transaction", round(k["avg_transaction"], 2)
    for c in ("B5", "B6", "B7", "B11"):
        ws[c].number_format = '"$"#,##0.00' if c == "B11" else '"$"#,##0'
    ws["B8"].number_format = "0.0%"
    ws["B9"].number_format = ws["B10"].number_format = "#,##0"

    ws["A13"], ws["A13"].font = "Category lookup (pick a category in B14)", Font(bold=True, color="0B1F3A")
    ws["A14"], ws["B14"] = "Category", cats["category"].iloc[0]
    dv = DataValidation(type="list", formula1=f"=Categories!$A$2:$A${cat_last}", allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B14")
    ws["A15"], ws["B15"] = "Spend (INDEX/MATCH)", f"=INDEX(Categories!$B$2:$B${cat_last},MATCH(B14,Categories!$A$2:$A${cat_last},0))"
    ws["A16"], ws["B16"] = "Share (VLOOKUP)", f"=VLOOKUP(B14,Categories!$A$2:$C${cat_last},3,FALSE)"
    ws["A17"], ws["B17"] = "Rank by spend", f"=RANK(B15,Categories!$B$2:$B${cat_last})"
    ws["B15"].number_format, ws["B16"].number_format = '"$"#,##0', "0.0%"
    ws["A19"], ws["A19"].font = "Key takeaways (computed)", Font(bold=True, color="0B1F3A")
    for i, t in enumerate(_takeaways(ctx), 20):
        ws.cell(i, 1, f"- {t}")
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 22

    # Forecast
    wf = wb.create_sheet("Forecast")
    fc = ctx["forecast"]
    wf["A1"], wf["A1"].font = f"Forecast - selected model: {fc.get('best_model') or fc['method']}", Font(bold=True, color="0B1F3A")
    header(wf, 3, ["Month", "Forecast", "Low (10%)", "High (90%)"])
    for r, row in enumerate(fc["future"].itertuples(), 4):
        wf.cell(r, 1, row.month.to_pydatetime()).number_format = "mmm yyyy"
        wf.cell(r, 2, float(row.forecast)).number_format = '"$"#,##0'
        wf.cell(r, 3, float(getattr(row, "lower", row.forecast))).number_format = '"$"#,##0'
        wf.cell(r, 4, float(getattr(row, "upper", row.forecast))).number_format = '"$"#,##0'
    r0 = 5 + len(fc["future"])
    if fc["metrics"]:
        wf.cell(r0, 1, f"Rolling-origin backtest, {fc['holdout_months']} held-out months (lower is better)").font = Font(bold=True)
        header(wf, r0 + 1, ["Model", "MAE", "RMSE", "MAPE %", "Best"])
        for i, m in enumerate(fc["metrics"], r0 + 2):
            for c, v in enumerate([m["model"], m["MAE"], m["RMSE"], m["MAPE_pct"], "yes" if m["best"] else ""], 1):
                wf.cell(i, c, v)
        r0 = r0 + 3 + len(fc["metrics"])
        header(wf, r0, ["Month", "Actual", "Random forest", "XGBoost", "ETS", "ARIMA", "Last month", "Seasonal naive"])
        for i, b in enumerate(fc["backtest"].itertuples(), r0 + 1):
            vals = [b.month.to_pydatetime(), b.actual, b.model, b.xgboost, b.ets, b.arima, b.last_month, b.seasonal_naive]
            for c, v in enumerate(vals, 1):
                cell = wf.cell(i, c, None if pd.isna(v) else v)
                cell.number_format = "mmm yyyy" if c == 1 else '"$"#,##0'
    wf.column_dimensions["A"].width = 30
    for col in "BCDEFGH":
        wf.column_dimensions[col].width = 16

    # Segments
    wg = wb.create_sheet("Segments")
    header(wg, 1, ["Segment", "Customers", "Avg days since last", "Avg transactions", "Avg spend", "Total spend"])
    for r, s in enumerate(ctx["segments"].itertuples(), 2):
        for c, v in enumerate([s.segment, s.customers, s.avg_recency_days, s.avg_transactions, s.avg_spend, s.total_spend], 1):
            wg.cell(r, c, v)
    for col in "ABCDEF":
        wg.column_dimensions[col].width = 20

    # Merchants
    wmr = wb.create_sheet("Merchants")
    mr = ctx["merchants"]
    if "error" in mr:
        wmr["A1"] = mr["error"]
    else:
        header(wmr, 1, ["Merchant", "Category", "Last month", "Forecast next month", "Change %", "Trailing 12m"])
        for r, t in enumerate(mr["table"].itertuples(), 2):
            for c, v in enumerate([t.merchant, t.category, t.last_month_spend, t.forecast_next_month, t.change_pct, t.trailing_12m_spend], 1):
                wmr.cell(r, c, v)
        n = len(mr["table"]) + 3
        wmr.cell(n, 1, f"Backtest ({mr['backtest_months']} months) WAPE %:").font = Font(bold=True)
        for i, m in enumerate(mr["metrics"], n + 1):
            wmr.cell(i, 1, m["model"])
            wmr.cell(i, 2, m["WAPE_pct"])
    for col in "ABCDEF":
        wmr.column_dimensions[col].width = 24

    # Notes
    wn = wb.create_sheet("Notes")
    notes = [
        "Methodology",
        "Forecasts: rolling-origin backtest; every prediction uses only earlier months.",
        "Models compared: random forest, XGBoost, Holt-Winters ETS, ARIMA, last-month and seasonal-naive baselines.",
        "The selected model is the lowest-MAE model on that same backtest, so its error is mildly optimistic.",
        "Prediction band: 10th-90th percentile of backtest errors, widened with horizon.",
        "Formulas on the Summary and Monthly sheets are live Excel formulas (SUM, AVERAGE, INDEX/MATCH, VLOOKUP, RANK).",
        "",
        DISCLAIMER,
    ]
    for i, t in enumerate(notes, 1):
        wn.cell(i, 1, t)
    wn["A1"].font = Font(bold=True, size=14, color="0B1F3A")
    wn.column_dimensions["A"].width = 120

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------------------------------ PowerPoint
def build_pptx(ctx: dict) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Emu, Inches, Pt

    NAVY, BLUE, GOLD, GREY = RGBColor(0x0B, 0x1F, 0x3A), RGBColor(0x1B, 0x3B, 0x69), RGBColor(0xB8, 0x89, 0x3A), RGBColor(0x55, 0x65, 0x7A)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def textbox(slide, x, y, w, h, text, size=16, bold=False, color=NAVY):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        lines = text if isinstance(text, list) else [text]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, color
            p.space_after = Pt(8)
        return tb

    def new_slide(title, subtitle=None):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = NAVY
        bar.line.fill.background()
        textbox(s, 0.6, 0.35, 12, 0.8, title, 28, True)
        if subtitle:
            textbox(s, 0.6, 1.05, 12, 0.5, subtitle, 14, False, GREY)
        textbox(s, 0.6, 7.0, 12, 0.4, DISCLAIMER, 9, False, GREY)
        return s

    def table(slide, x, y, w, rows, col_widths=None, font=12):
        shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(0.4 * len(rows)))
        t = shape.table
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cell = t.cell(r, c)
                cell.text = str(val)
                para = cell.text_frame.paragraphs[0]
                para.font.size = Pt(font)
                para.font.bold = r == 0
                if r == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = NAVY
                    para.font.color.rgb = RGBColor(255, 255, 255)
        if col_widths:
            for i, cw in enumerate(col_widths):
                t.columns[i].width = Inches(cw)
        return t

    k, fc = ctx["kpis"], ctx["forecast"]
    # 1 Title
    s = prs.slides.add_slide(blank)
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.fill.background()
    textbox(s, 0.9, 2.3, 11.5, 1.2, "SpendIQ", 54, True, RGBColor(255, 255, 255))
    textbox(s, 0.9, 3.6, 11.5, 1.0, "Spend analytics and forecast briefing", 26, False, RGBColor(0xDB, 0xE5, 0xF3))
    textbox(s, 0.9, 4.6, 11.5, 0.8, f"{ctx['label']}  |  {ctx['period']}  |  generated {ctx['generated']}", 14, False, RGBColor(0xB6, 0xC8, 0xE4))
    textbox(s, 0.9, 6.5, 11.5, 0.6, DISCLAIMER, 10, False, RGBColor(0xB6, 0xC8, 0xE4))

    # 2 Executive summary
    s = new_slide("Executive summary", "Headline figures for the selected scope")
    tiles = [("Total spend", f"${k['total_spend']:,.0f}"), ("Transactions", f"{k['transactions']:,}"),
             ("Customers", f"{k['customers']:,}"), ("Avg transaction", f"${k['avg_transaction']:,.2f}")]
    for i, (lab, val) in enumerate(tiles):
        x = 0.6 + i * 3.1
        box = s.shapes.add_shape(1, Inches(x), Inches(1.8), Inches(2.9), Inches(1.5))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(0xEE, 0xF3, 0xFA)
        box.line.color.rgb = RGBColor(0xDB, 0xE5, 0xF3)
        textbox(s, x + 0.15, 1.9, 2.6, 0.4, lab, 12, False, GREY)
        textbox(s, x + 0.15, 2.35, 2.6, 0.8, val, 26, True)
    textbox(s, 0.6, 3.8, 12, 3, _takeaways(ctx), 16)

    # 3 Monthly chart
    s = new_slide("Monthly spend", "Total transaction value by month")
    cd = CategoryChartData()
    cd.categories = [f"{m:%b %y}" for m in ctx["monthly"]["month"]]
    cd.add_series("Spend", [round(float(v), 2) for v in ctx["monthly"]["spend"]])
    ch = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(0.6), Inches(1.6), Inches(12), Inches(5.2), cd).chart
    ch.has_legend = False
    ch.series[0].format.line.color.rgb = BLUE

    # 4 Categories
    s = new_slide("Spend by merchant category")
    cd = CategoryChartData()
    cats = ctx["categories"]
    cd.categories = list(cats["category"])
    cd.add_series("Spend", [round(float(v), 2) for v in cats["spend"]])
    ch = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(1.3), Inches(12), Inches(5.5), cd).chart
    ch.has_legend = False
    ch.series[0].format.fill.solid()
    ch.series[0].format.fill.fore_color.rgb = BLUE

    # 5 Forecast
    s = new_slide("Spend forecast", f"{fc.get('best_model') or fc['method']} | next {len(fc['future'])} months")
    hist = ctx["monthly"].tail(18)
    cd = CategoryChartData()
    labels = [f"{m:%b %y}" for m in hist["month"]] + [f"{m:%b %y}" for m in fc["future"]["month"]]
    cd.categories = labels
    cd.add_series("Actual", [round(float(v), 2) for v in hist["spend"]] + [None] * len(fc["future"]))
    cd.add_series("Forecast", [None] * (len(hist) - 1) + [round(float(hist["spend"].iloc[-1]), 2)] + [float(v) for v in fc["future"]["forecast"]])
    ch = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.6), Inches(1.6), Inches(12), Inches(5.2), cd).chart
    ch.has_legend, ch.legend.position = True, XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.series[0].format.line.color.rgb = BLUE
    ch.series[1].format.line.color.rgb = GOLD

    # 6 Model comparison
    s = new_slide("How the forecast was validated", "Rolling-origin backtest - each prediction uses only earlier months")
    if fc["metrics"]:
        rows = [["Model", "MAE", "RMSE", "MAPE", ""]] + [
            [m["model"], f"${m['MAE']:,.0f}", f"${m['RMSE']:,.0f}", f"{m['MAPE_pct']:.1f}%", "best" if m["best"] else ""]
            for m in fc["metrics"]]
        table(s, 0.6, 1.8, 12, rows, [5.6, 1.7, 1.7, 1.5, 1.5])
        textbox(s, 0.6, 5.0, 12, 1.5, [
            f"Evaluation: {fc['holdout_months']} held-out months, one-step-ahead.",
            "Baselines (last-month, seasonal-naive) are included so model value is measured, not assumed.",
            "The selected model is chosen on this same backtest, so its reported error is mildly optimistic."], 14)
    else:
        textbox(s, 0.6, 1.8, 12, 1.5, "Not enough history in this slice for a backtest; a flat last-month forecast is shown.", 16)

    # 7 Segments
    s = new_slide("Customer segments", "Log-scaled RFM + KMeans; descriptive labels, not value judgments")
    seg = ctx["segments"]
    if len(seg):
        rows = [["Segment", "Customers", "Avg days since last", "Avg txns", "Avg spend", "Total spend"]] + [
            [r.segment, f"{r.customers:,}", f"{r.avg_recency_days:.0f}", f"{r.avg_transactions:.1f}", f"${r.avg_spend:,.0f}", f"${r.total_spend:,.0f}"]
            for r in seg.itertuples()]
        table(s, 0.6, 1.8, 12, rows)

    # 8 Merchants
    s = new_slide("Merchant outlook", "Pooled XGBoost one-step-ahead forecast per merchant")
    mr = ctx["merchants"]
    if "error" in mr:
        textbox(s, 0.6, 1.8, 12, 1, mr["error"], 16)
    else:
        top = mr["table"].head(8)
        rows = [["Merchant", "Category", "Last month", "Next month (fcst)", "Change"]] + [
            [r.merchant, r.category, f"${r.last_month_spend:,.0f}", f"${r.forecast_next_month:,.0f}", f"{r.change_pct:+.1f}%"]
            for r in top.itertuples()]
        table(s, 0.6, 1.8, 12, rows, font=12)
        best = min(mr["metrics"], key=lambda m: m["WAPE_pct"])
        textbox(s, 0.6, 5.6, 12, 1, f"Backtest over {mr['backtest_months']} months: best WAPE {best['WAPE_pct']:.1f}% ({best['model']}).", 14)

    # 9 Limitations
    s = new_slide("Method, limitations and responsible use")
    textbox(s, 0.6, 1.4, 12, 5.4, [
        "Data is synthetic and generated locally; seasonality and response effects are built in by construction.",
        "Model performance on synthetic data does not establish real-world performance.",
        "The campaign model is transaction-level and is not a causal uplift model.",
        "Fairness, drift and privacy reviews are not comprehensive in this demonstration.",
        "Human, legal and compliance review would be required before any real-world use.",
        "Out of scope: credit or lending, eligibility or access to financial services, fraud decisions, pricing, real-customer targeting.",
    ], 16)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
