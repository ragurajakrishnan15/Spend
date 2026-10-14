"""Tests for merchant cleaning, extended forecasting models, merchant forecasts and exports."""
import io
import sys
from pathlib import Path

import duckdb
import numpy as np
import pandas as pd
import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.analytics import forecast_merchant_spend, forecast_monthly_spend, monthly_spend, run_sql  # noqa: E402
from src.data import generate_synthetic_data  # noqa: E402
from src.exports import build_context, build_excel, build_pptx  # noqa: E402
from src.merchants import (CATALOG, UNKNOWN, classification_report, classify_descriptors,  # noqa: E402
                           normalize_descriptor)


@pytest.fixture(scope="module")
def df():
    d = generate_synthetic_data(seed=11, customers=200, transactions=9000)
    return d.join(classify_descriptors(d["merchant_descriptor"])[["merchant_clean", "category_clean"]])


# ------------------------------------------------------------------ merchant regex rules
@pytest.mark.parametrize("raw,expected", [
    ("POS DEBIT GREENFIELD MKT #4821 NEW YORK NY", "Greenfield Market"),
    ("SQ *BLUE FIG BISTRO", "Blue Fig Bistro"),
    ("TST* CPR KTL CAFE #77 AUSTIN TX", "Copper Kettle Cafe"),
    ("PURCHASE AUTHORIZED ON 0412 STARLIGHT CINEMAS DENVER CO", "Starlight Cinemas"),
    ("SMMT TELECOM", "Summit Telecom"),
    ("random vendor llc", UNKNOWN),
    ("", UNKNOWN),
])
def test_regex_rules_known_descriptors(raw, expected):
    assert classify_descriptors(pd.Series([raw])).iloc[0]["merchant_clean"] == expected


def test_normalize_strips_noise():
    assert normalize_descriptor("SQ *Greenfield Market #123 Boston MA") == "GREENFIELD MARKET BOSTON"


def test_every_catalog_variant_matches_its_own_rule():
    for m in CATALOG:
        out = classify_descriptors(pd.Series(list(m.variants)))
        assert (out["merchant_clean"] == m.name).all(), m.name


def test_classification_report_quality(df):
    rep = classification_report(df, df[["merchant_clean", "category_clean", "merchant_descriptor"]].assign(
        normalized=df["merchant_descriptor"].map(normalize_descriptor)))
    assert rep["coverage"] > 0.85          # ~7% garbled by design
    assert rep["precision"] > 0.98
    assert rep["category_agreement"] > 0.98
    assert len(rep["top_unmatched"]) > 0


def test_generator_has_garbled_unmatched_share(df):
    share = (df["merchant_clean"] == UNKNOWN).mean()
    assert 0.02 < share < 0.15


# ------------------------------------------------------------------ extended forecast models
def test_forecast_benchmarks_all_models():
    d = generate_synthetic_data(seed=5, customers=100, transactions=6000)
    res = forecast_monthly_spend(monthly_spend(d), 3)
    keys = {m["key"] for m in res["metrics"]}
    assert keys == {"model", "xgboost", "ets", "arima", "last_month", "seasonal_naive"}
    assert sum(m["best"] for m in res["metrics"]) == 1
    assert res["best_model"] == min(res["metrics"], key=lambda r: r["MAE"])["model"]
    assert len(res["future"]) == 3
    assert set(res["future_by_model"]) >= {"XGBoost (lag features)", "Holt-Winters ETS", "Seasonal naive (t-12)"}
    assert all(len(v) == 3 for v in res["future_by_model"].values())
    assert (res["future"]["forecast"] > 0).all()


def test_statistical_models_degrade_gracefully_on_short_history():
    m = pd.DataFrame({"month": pd.date_range("2024-01-01", periods=10, freq="MS"),
                      "spend": np.linspace(100, 160, 10) + np.tile([0, 5], 5)})
    res = forecast_monthly_spend(m, 2)
    assert res["metrics"]
    assert "seasonal_naive" not in {x["key"] for x in res["metrics"]}


# ------------------------------------------------------------------ merchant forecasting
def test_merchant_forecast(df):
    res = forecast_merchant_spend(df)
    assert "error" not in res
    assert {m["model"] for m in res["metrics"]} >= {"Last-month naive", "Seasonal naive (t-12)"}
    t = res["table"]
    assert t["merchant"].nunique() == len(t) and (t["forecast_next_month"] >= 0).all()
    assert UNKNOWN not in set(t["merchant"])
    assert res["next_month"] == pd.Timestamp("2026-01-01")


def test_merchant_forecast_edge_cases(df):
    assert "error" in forecast_merchant_spend(df.iloc[0:0])
    short = df[df["timestamp"] < "2023-09-01"]
    assert "error" in forecast_merchant_spend(short)


# ------------------------------------------------------------------ exports
def test_excel_export_has_live_formulas(df):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(build_excel(build_context(df))))
    assert {"Summary", "Monthly", "Categories", "Forecast", "Segments", "Merchants", "Notes"} <= set(wb.sheetnames)
    s = wb["Summary"]
    formulas = [str(c.value) for row in s.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert any("INDEX(" in f and "MATCH(" in f for f in formulas)
    assert any(f.startswith("=VLOOKUP(") for f in formulas)
    assert any(f.startswith("=SUM(") for f in formulas)
    assert wb["Monthly"]["C3"].value == "=B3/B2-1"
    assert len(wb["Monthly"]._charts) == 1


def test_pptx_export_structure(df):
    from pptx import Presentation
    prs = Presentation(io.BytesIO(build_pptx(build_context(df))))
    assert len(prs.slides) == 9
    kinds = [sh.has_chart for sl in prs.slides for sh in sl.shapes]
    assert sum(kinds) == 3
    text = " ".join(sh.text_frame.text for sl in prs.slides for sh in sl.shapes if sh.has_text_frame)
    assert "Synthetic data only" in text


def test_export_rejects_empty_scope(df):
    with pytest.raises(ValueError):
        build_context(df.iloc[0:0])


# ------------------------------------------------------------------ SQL regex support + API
def test_sql_regex_functions_and_cleaned_merchant_column(tmp_path, df):
    path = tmp_path / "r.duckdb"
    con = duckdb.connect(str(path))
    con.register("v", df)
    con.execute("CREATE TABLE transactions AS SELECT * FROM v")
    con.close()
    out = run_sql("SELECT merchant_clean, COUNT(*) n FROM transactions "
                  "WHERE regexp_matches(upper(merchant_descriptor), 'GREEN\\s?F(IE)?LD|GRNFLD') GROUP BY 1", str(path))
    assert list(out["merchant_clean"]) == ["Greenfield Market"]


@pytest.fixture(scope="module")
def client():
    from backend.main import app
    return TestClient(app)


def test_api_merchants(client):
    r = client.get("/api/merchants").json()
    assert r["classification"]["coverage"] > 0.85
    assert r["forecast"]["table"] and r["forecast"]["metrics"]


def test_api_classify(client):
    r = client.post("/api/merchants/classify", json={"descriptor": "SQ *CLRSTRM WATER #12 BOSTON MA"}).json()
    assert r["matched"] and r["merchant"] == "ClearStream Water" and r["category"] == "Utilities"
    r = client.post("/api/merchants/classify", json={"descriptor": "ZZZ UNKNOWN SHOP"}).json()
    assert not r["matched"] and r["merchant"] == "Unknown"


def test_api_exports(client):
    x = client.get("/api/export/excel", params={"category": "Travel"})
    assert x.status_code == 200 and x.content[:2] == b"PK" and "spreadsheetml" in x.headers["content-type"]
    p = client.get("/api/export/pptx")
    assert p.status_code == 200 and p.content[:2] == b"PK" and "presentationml" in p.headers["content-type"]
    assert client.get("/api/export/excel", params={"start": "2099-01-01"}).status_code == 404


def test_api_regex_sql_examples_execute(client):
    for ex in client.get("/api/sql/examples").json()["examples"]:
        r = client.post("/api/sql", json={"query": ex["query"]})
        assert r.status_code == 200, (ex["name"], r.text)


# ------------------------------------------------------------------ operations, budget, macro, narrative
from src import narrative  # noqa: E402
from src.analytics import DEFAULT_BUDGETS, budget_variance, run_operations_forecast  # noqa: E402
from src.data import CATEGORIES  # noqa: E402
from src.macro_intelligence import get_all_indicators, get_macro_context, get_report_library  # noqa: E402


def test_operations_forecast_returns_required_keys(df):
    res = run_operations_forecast(df, horizon=3)
    assert {"transaction_count", "avg_ticket", "channels", "horizon", "aggregate_metrics"} <= res.keys()
    assert res["horizon"] == 3
    series = res["transaction_count"][0]
    assert {"channel", "history", "future", "xgb_mae", "naive_mae", "best_model"} <= series.keys()
    assert len(series["future"]) == 3
    assert all(f["value"] >= 0 for ch in res["avg_ticket"] for f in ch["future"])


def test_operations_forecast_empty_frame(df):
    assert "error" in run_operations_forecast(df.iloc[0:0])


def test_budget_variance_rag_status(df):
    res = budget_variance(df)
    assert {r["rag"] for r in res["detail"]} <= {"green", "amber", "red"}
    for r in res["detail"]:
        assert r["rag"] == ("green" if r["actual"] <= r["budget"]
                            else "amber" if r["variance_pct"] <= 10 else "red")


def test_default_budgets_cover_every_generated_category():
    assert set(DEFAULT_BUDGETS) == set(CATEGORIES)


def test_budget_override_changes_status(df):
    tiny = {c: 1.0 for c in CATEGORIES}
    assert {r["rag"] for r in budget_variance(df, tiny)["detail"]} == {"red"}


def test_macro_context_keys():
    ctx = get_macro_context()
    assert {"month", "indicators", "themes", "watch_items"} <= ctx.keys()
    assert {"cpi_yoy", "consumer_sentiment", "gdp_growth"} <= ctx["indicators"].keys()
    assert ctx["themes"]
    assert get_macro_context("1999-01")["month"] == get_all_indicators()[-1]["month"]  # unknown month -> latest


def test_report_library_tags_match_generated_categories():
    reports = get_report_library()
    assert len(reports) == 6
    tags = {t for r in reports for t in r["category_tags"]}
    assert tags <= set(CATEGORIES)
    assert all(r["category_tags"] for r in get_report_library("Travel"))


def test_narrative_fallback_no_api_key(monkeypatch):
    monkeypatch.setattr(narrative, "_GEMINI_KEY", "")
    for section in ["overview", "forecast", "segments", "campaign", "ops", "budget", "macro", "merchants"]:
        assert narrative.generate_summary(section, {}).strip()


def test_narrative_falls_back_when_gemini_unavailable(monkeypatch):
    monkeypatch.setattr(narrative, "_GEMINI_KEY", "fake-key")
    out = narrative.generate_summary("ops", {})
    assert out == narrative._template_summary("ops", {})


def test_api_ops_budget_macro_narrative(client):
    ops = client.get("/api/ops-forecast", params={"horizon": 2}).json()
    assert ops["horizon"] == 2 and ops["channels"]
    assert client.get("/api/ops-forecast", params={"horizon": 99}).status_code == 422
    b = client.get("/api/budget").json()
    assert b["summary"] and {r["rag"] for r in b["detail"]} <= {"green", "amber", "red"}
    m = client.get("/api/macro").json()
    assert m["context"]["watch_items"] is not None and m["indicators"]
    lib = client.get("/api/report-library", params={"category": "Travel"}).json()
    assert lib["total"] == len(lib["reports"]) >= 1
    n = client.post("/api/narrative/budget", json={"context": {"categories_over_budget": 3, "total_categories": 8}})
    assert n.status_code == 200 and n.json()["summary"]
    assert client.post("/api/narrative/nope", json={}).status_code == 400


def test_every_sql_example_runs_and_has_warehouse_translations(client):
    ex = client.get("/api/sql/examples").json()["examples"]
    assert len(ex) >= 12
    for e in ex:
        out = client.post("/api/sql", json={"query": e["query"]})
        assert out.status_code == 200, e["name"]
        assert e["snowflake"].strip() and e["athena"].strip(), e["name"]
    names = {e["name"] for e in ex}
    assert {"Budget variance with RAG status", "Operations: volume by channel",
            "Quarterly spend by category", "Campaign response by exposure"} <= names


def test_sql_budget_example_matches_python_budget_variance(client, df):
    ex = {e["name"]: e for e in client.get("/api/sql/examples").json()["examples"]}
    sql = client.post("/api/sql", json={"query": ex["Budget variance with RAG status"]["query"]}).json()
    api = client.get("/api/budget").json()["detail"]
    assert sql["row_count"] == len(api)
    cols = sql["columns"]
    rag_sql = {(r[cols.index("month")][:7], r[cols.index("category")]): r[cols.index("rag")] for r in sql["rows"]}
    rag_py = {(r["month"], r["category"]): r["rag"] for r in api}
    assert rag_sql == rag_py
